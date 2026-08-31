from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


def generate_office_files(root: Path) -> dict[str, Path]:
    root.mkdir(parents=True, exist_ok=True)

    docx_path = root / "sample.docx"
    document = Document()
    document.add_heading("Parser evaluation", level=1)
    document.add_paragraph("DOCX fixture")
    document.save(docx_path)

    pptx_path = root / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Parser evaluation"
    presentation.save(pptx_path)

    xlsx_path = root / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evaluation"
    sheet.append(["parser", "score"])
    sheet.append(["fixture", 1])
    workbook.save(xlsx_path)

    return {"docx": docx_path, "pptx": pptx_path, "xlsx": xlsx_path}
